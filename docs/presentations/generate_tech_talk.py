#!/usr/bin/env python3
"""Generate NetApp GSD Recipe tech-talk deck (one idea, minimal text, high contrast)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Palette — teal trust + dark sandwich
NAVY = RGBColor(0x21, 0x29, 0x5C)
TEAL = RGBColor(0x02, 0x80, 0x90)
MINT = RGBColor(0x02, 0xC3, 0x9A)
CORAL = RGBColor(0xF9, 0x61, 0x67)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xF4, 0xF6, 0xF8)
CHARCOAL = RGBColor(0x36, 0x45, 0x4F)
SLATE = RGBColor(0x64, 0x74, 0x8B)

OUT = Path(__file__).resolve().parent / "NetApp-GSD-Recipe-Tech-Talk.pptx"
VIDEO = Path(__file__).resolve().parent / "assets" / "recipe-demo-post-usage.mp4"


def set_slide_bg(slide, rgb: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    size=28,
    bold=False,
    color=CHARCOAL,
    align=PP_ALIGN.LEFT,
    font="Calibri",
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = color
    return box


def add_accent_bar(slide, left, top, height, color=TEAL, width=Inches(0.12)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def slide_title_dark(prs, title: str, subtitle: str = "") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, NAVY)
    add_textbox(
        slide,
        Inches(0.9),
        Inches(2.0),
        Inches(11.5),
        Inches(1.4),
        title,
        size=44,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.LEFT,
        font="Calibri Light",
    )
    if subtitle:
        add_textbox(
            slide,
            Inches(0.9),
            Inches(3.5),
            Inches(11.5),
            Inches(0.8),
            subtitle,
            size=22,
            color=MINT,
            align=PP_ALIGN.LEFT,
        )
    add_accent_bar(slide, Inches(0.9), Inches(1.7), Inches(0.08), MINT)


def slide_one_idea(prs, headline: str, sub: str, bg=OFF_WHITE) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, bg)
    add_accent_bar(slide, Inches(0.7), Inches(1.2), Inches(4.5), TEAL)
    add_textbox(
        slide,
        Inches(1.0),
        Inches(1.5),
        Inches(11.0),
        Inches(2.2),
        headline,
        size=40,
        bold=True,
        color=NAVY,
    )
    add_textbox(
        slide,
        Inches(1.0),
        Inches(3.8),
        Inches(10.5),
        Inches(1.2),
        sub,
        size=22,
        color=SLATE,
    )


def slide_prerequisites(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, OFF_WHITE)
    add_textbox(
        slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.7),
        "Before you start", size=32, bold=True, color=NAVY,
    )

    cols = [
        ("Must have", TEAL, [
            "Cursor + Agent",
            "Git repo (target)",
            "python3 + git",
            "Atlassian MCP (Jira)",
            "gh auth (GitHub PR)",
        ]),
        ("Soft / optional", SLATE, [
            "GSD core (npx)",
            "graphify + uv",
            "recipe-validate-tokens first",
        ]),
    ]
    x = Inches(0.8)
    for title, accent, items in cols:
        box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, x, Inches(1.3), Inches(5.8), Inches(4.8)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = accent
        add_textbox(
            slide, x + Inches(0.25), Inches(1.55), Inches(5.3), Inches(0.45),
            title, size=22, bold=True, color=accent,
        )
        y = Inches(2.2)
        for item in items:
            add_textbox(
                slide, x + Inches(0.35), y, Inches(5.1), Inches(0.4),
                item, size=18, color=CHARCOAL,
            )
            y += Inches(0.55)
        x += Inches(6.2)

    add_textbox(
        slide, Inches(0.8), Inches(6.2), Inches(11.5), Inches(0.5),
        "install.sh --verify  ·  recipe-validate-tokens  ·  see README Prerequisites",
        size=14, color=SLATE,
    )


def slide_big_stat(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, OFF_WHITE)

    # Ad-hoc column
    box1 = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.8), Inches(4.8), Inches(3.8)
    )
    box1.fill.solid()
    box1.fill.fore_color.rgb = RGBColor(0xFE, 0xE2, 0xE2)
    box1.line.color.rgb = CORAL
    add_textbox(
        slide,
        Inches(1.3),
        Inches(2.2),
        Inches(4.2),
        Inches(0.5),
        "Ad-hoc AI coding",
        size=20,
        bold=True,
        color=CORAL,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        Inches(1.3),
        Inches(3.0),
        Inches(4.2),
        Inches(1.2),
        "~2 days",
        size=56,
        bold=True,
        color=CORAL,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        Inches(1.3),
        Inches(4.5),
        Inches(4.2),
        Inches(0.6),
        "KB-Evaluations · AgentStudio",
        size=14,
        color=SLATE,
        align=PP_ALIGN.CENTER,
    )

    # Arrow
    add_textbox(
        slide,
        Inches(5.9),
        Inches(3.2),
        Inches(1.0),
        Inches(0.8),
        "→",
        size=48,
        bold=True,
        color=TEAL,
        align=PP_ALIGN.CENTER,
    )

    # Recipe column
    box2 = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(7.2), Inches(1.8), Inches(4.8), Inches(3.8)
    )
    box2.fill.solid()
    box2.fill.fore_color.rgb = RGBColor(0xD1, 0xFA, 0xE5)
    box2.line.color.rgb = MINT
    add_textbox(
        slide,
        Inches(7.5),
        Inches(2.2),
        Inches(4.2),
        Inches(0.5),
        "NetApp GSD Recipe",
        size=20,
        bold=True,
        color=TEAL,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        Inches(7.5),
        Inches(3.0),
        Inches(4.2),
        Inches(1.2),
        "3–6 hours",
        size=56,
        bold=True,
        color=TEAL,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        Inches(7.5),
        Inches(4.5),
        Inches(4.2),
        Inches(0.6),
        "KAN-53 · PR #465",
        size=14,
        color=SLATE,
        align=PP_ALIGN.CENTER,
    )

    add_textbox(
        slide,
        Inches(1.0),
        Inches(0.6),
        Inches(11.0),
        Inches(0.6),
        "Same feature. Same engineer. Different workflow.",
        size=28,
        bold=True,
        color=NAVY,
    )


def slide_flow(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, OFF_WHITE)
    add_textbox(
        slide,
        Inches(0.8),
        Inches(0.5),
        Inches(11.0),
        Inches(0.7),
        "One runway: PRD → Jira → Ship",
        size=32,
        bold=True,
        color=NAVY,
    )

    steps = [
        ("Install", "1 bash cmd"),
        ("Onboard", "PRD + Epic"),
        ("Plan & Run", "GSD phases"),
        ("Verify & PR", "Ship gate"),
        ("Jira sync", "Traceability"),
    ]
    x = Inches(0.5)
    w = Inches(2.2)
    for i, (title, sub) in enumerate(steps):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(2.0), w, Inches(2.8))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE if i % 2 == 0 else RGBColor(0xE0, 0xF2, 0xF1)
        shape.line.color.rgb = TEAL
        add_textbox(
            slide, x + Inches(0.15), Inches(2.4), w - Inches(0.3), Inches(0.6),
            title, size=22, bold=True, color=TEAL, align=PP_ALIGN.CENTER,
        )
        add_textbox(
            slide, x + Inches(0.15), Inches(3.2), w - Inches(0.3), Inches(0.8),
            sub, size=16, color=SLATE, align=PP_ALIGN.CENTER,
        )
        if i < len(steps) - 1:
            add_textbox(
                slide, x + w + Inches(0.05), Inches(2.9), Inches(0.35), Inches(0.5),
                "→", size=28, bold=True, color=TEAL, align=PP_ALIGN.CENTER,
            )
        x += w + Inches(0.4)


def slide_code(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, OFF_WHITE)
    add_textbox(
        slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.7),
        "Install once. Then Cursor only.", size=32, bold=True, color=NAVY,
    )

    code_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.5), Inches(2.2)
    )
    code_box.fill.solid()
    code_box.fill.fore_color.rgb = RGBColor(0x1E, 0x1E, 0x1E)
    code_box.line.fill.background()

    lines = [
        "./bench/runners/install-recipe-to-target.sh \\",
        "  --target /<path-to-repo> --verify",
    ]
    tf = slide.shapes.add_textbox(
        Inches(1.1), Inches(2.1), Inches(11.0), Inches(1.8)
    ).text_frame
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Consolas"
        p.font.size = Pt(22)
        p.font.color.rgb = MINT

    add_textbox(
        slide, Inches(0.8), Inches(4.3), Inches(11), Inches(1.0),
        "Open target repo in Cursor → invoke recipe-* by name → recipe-help for the catalog",
        size=20, color=SLATE,
    )


def slide_commands(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, OFF_WHITE)
    add_textbox(
        slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.7),
        "Five commands you'll actually run", size=32, bold=True, color=NAVY,
    )

    rows = [
        ("recipe-validate-tokens", "Preflight GitHub + Jira"),
        ("recipe-onboard @PRD.md", "Epic + phase tasks"),
        ("recipe-bootstrap-knowledge", "Map + graphify context"),
        ("recipe-plan-phase N → run-phase N", "Plan AND run (same phase)"),
        ("recipe-run-phases 2 5 --full", "OR loop verify→PR→settle"),
    ]
    y = Inches(1.3)
    for cmd, desc in rows:
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), y, Inches(0.08), Inches(0.85)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = TEAL
        bar.line.fill.background()
        add_textbox(
            slide, Inches(1.05), y + Inches(0.05), Inches(5.5), Inches(0.45),
            cmd, size=18, bold=True, color=TEAL, font="Consolas",
        )
        add_textbox(
            slide, Inches(6.8), y + Inches(0.1), Inches(5.5), Inches(0.45),
            desc, size=18, color=CHARCOAL,
        )
        y += Inches(1.05)


def slide_and_or(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, OFF_WHITE)
    add_textbox(
        slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.7),
        "Don't double-run phases", size=32, bold=True, color=NAVY,
    )

    # AND box
    and_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.5), Inches(2.5)
    )
    and_box.fill.solid()
    and_box.fill.fore_color.rgb = WHITE
    and_box.line.color.rgb = TEAL
    add_textbox(
        slide, Inches(1.1), Inches(1.9), Inches(5.0), Inches(0.5),
        "AND — one phase", size=24, bold=True, color=TEAL,
    )
    add_textbox(
        slide, Inches(1.1), Inches(2.6), Inches(5.0), Inches(1.2),
        "plan-phase N\nthen run-phase N",
        size=20, color=CHARCOAL, font="Consolas",
    )

    # OR box
    or_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.5), Inches(2.5)
    )
    or_box.fill.solid()
    or_box.fill.fore_color.rgb = WHITE
    or_box.line.color.rgb = MINT
    add_textbox(
        slide, Inches(7.1), Inches(1.9), Inches(5.0), Inches(0.5),
        "OR — phase range", size=24, bold=True, color=TEAL,
    )
    add_textbox(
        slide, Inches(7.1), Inches(2.6), Inches(5.0), Inches(1.2),
        "run-phases 2 5 --full\n(instead of manual 2…5)",
        size=20, color=CHARCOAL, font="Consolas",
    )

    add_textbox(
        slide, Inches(0.8), Inches(4.5), Inches(11.5), Inches(0.8),
        "Graphify: recipe-bootstrap-knowledge before first plan · /gsd-graphify query mid-phase",
        size=18, color=SLATE,
    )


def slide_chart(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, OFF_WHITE)
    add_textbox(
        slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.7),
        "Field pilot (directional, not SLA)", size=32, bold=True, color=NAVY,
    )

    from pptx.chart.data import ChartData

    data = ChartData()
    data.categories = ["Ad-hoc", "Recipe"]
    data.add_series("Wall-clock (hours)", (16, 4.5))

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(2.0), Inches(1.3), Inches(9.0), Inches(4.5),
        data,
    ).chart
    chart.has_legend = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.maximum_scale = 18

    add_textbox(
        slide, Inches(0.8), Inches(6.0), Inches(11.5), Inches(0.5),
        "KB-Evaluations · AgentStudio · Jul 2026 · github.com/vs06101996/NetApp-GSD-Recipe",
        size=14, color=SLATE,
    )


def slide_demo(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, NAVY)
    add_textbox(
        slide, Inches(0.9), Inches(0.5), Inches(5.8), Inches(0.9),
        "Post-recipe walkthrough", size=32, bold=True, color=WHITE,
    )
    add_textbox(
        slide, Inches(0.9), Inches(1.35), Inches(5.5), Inches(0.5),
        "~7 min · Jul 2026 recording", size=16, color=RGBColor(0xAA, 0xCC, 0xDD),
    )

    bullets = [
        "What changed after recipe-onboard",
        "Jira + PLAN.md in the repo",
        "Shipped output (PR / artifacts)",
        "Click video ▶ to play in Presenter View",
    ]
    y = Inches(2.1)
    for line in bullets:
        add_textbox(slide, Inches(1.0), y, Inches(5.2), Inches(0.45), line, size=18, color=MINT)
        y += Inches(0.55)

    if VIDEO.is_file():
        slide.shapes.add_movie(
            str(VIDEO),
            Inches(6.4),
            Inches(0.9),
            Inches(6.5),
            Inches(5.8),
            mime_type="video/mp4",
        )
    else:
        add_textbox(
            slide, Inches(6.4), Inches(2.5), Inches(6.2), Inches(2.0),
            "Place video at:\nassets/recipe-demo-post-usage.mp4\n\nThen re-run generate_tech_talk.py",
            size=16, color=RGBColor(0xAA, 0xCC, 0xDD),
        )

    add_textbox(
        slide, Inches(0.9), Inches(6.5), Inches(12), Inches(0.5),
        "Source: GMT20260717 post-usage screen recording (KB-Evaluations / AgentStudio context)",
        size=12, color=RGBColor(0x88, 0xAA, 0xBB),
    )


def slide_qa(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, NAVY)
    add_textbox(
        slide, Inches(0.9), Inches(2.5), Inches(11), Inches(1.2),
        "Questions?",
        size=54, bold=True, color=WHITE, align=PP_ALIGN.CENTER, font="Calibri Light",
    )
    add_textbox(
        slide, Inches(0.9), Inches(4.0), Inches(11), Inches(0.8),
        "NetApp-GSD-Recipe · recipe-help in Cursor",
        size=22, color=MINT, align=PP_ALIGN.CENTER,
    )


def add_notes(slide, text: str) -> None:
    notes = slide.notes_slide.notes_text_frame
    notes.text = text


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_title_dark(
        prs,
        "NetApp GSD Recipe",
        "Give AI a runway — not just a prompt",
    )
    add_notes(prs.slides[-1], "20-min talk. One message: structured recipe beats vibe coding. Skip long bio — 10 sec intro.")

    slide_prerequisites(prs)
    add_notes(prs.slides[-1], "Don't read every bullet. Hit: Cursor, Jira MCP, gh, then install. graphify optional.")

    slide_big_stat(prs)
    add_notes(prs.slides[-1], "Hook: real AgentStudio KB-Evaluations. Same you, same scope. Pause on the gap.")

    slide_one_idea(
        prs,
        "One idea to remember",
        "Wrap GSD with recipe-* commands: structured delivery + Jira traceability in Cursor.",
    )
    add_notes(prs.slides[-1], "GSD still orchestrates .planning/. Recipe adds commands + Jira. Not replacing GSD.")

    slide_flow(prs)
    add_notes(prs.slides[-1], "Walk left to right. Emphasize: install is bash once; everything else is Cursor.")

    slide_code(prs)
    add_notes(prs.slides[-1], "Live or screenshot. Point at --verify. No benchmark harness needed for delivery.")

    slide_commands(prs)
    add_notes(prs.slides[-1], "Don't read every line — pick 2-3. onboard + plan/run are the heart.")

    slide_and_or(prs)
    add_notes(prs.slides[-1], "Common mistake: run-phases AND manual plan-phase for same numbers. Pick one path.")

    slide_chart(prs)
    add_notes(prs.slides[-1], "Say caveats: N=1 pilot, directional. Harness N=5 still pending.")

    slide_demo(prs)
    add_notes(prs.slides[-1], "Play embedded 7-min video (post-usage + changes). Pause to highlight Jira/PR. Leave 5 min Q&A.")

    slide_qa(prs)
    add_notes(prs.slides[-1], "Backup: recipe-help vs gsd-help, graphify via bootstrap-knowledge, recipe-review-ship for PR.")

    prs.save(str(OUT))
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
